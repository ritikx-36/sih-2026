"""
Fill the OFFICIAL SIH Idea template with our Thermal-Shelter content.

Opens the official SIH 2026 Idea template (already titled "SMART INDIA HACKATHON
2026"), populates the six mandated sections (title page + 5 content slides), drops
our figures into the free zones, renders the References slide as a numbered list
with clickable links, deletes the "Important Instructions" slide, and saves an
editable, submission-ready .pptx. The template's own layout/branding/fonts are
preserved; only text, a few box sizes, and images change.

python-pptx lives on the SYSTEM interpreter here, so run:
    python3 scripts/fill_official_ppt.py
Figures must already exist in docs/ (they do).

Team Name (HimTaap) and Theme (Miscellaneous) are filled in. The SIH Team ID is
only allotted after internal selection, so the title page shows "To be allotted
after selection" for now — swap in the real portal Team ID once you have it.
"""
import copy
import os

from PIL import Image
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

SRC = "/Users/ritik/Downloads/SIH2026-IDEA-Presentation-Format.pptx"
OUT = "/Users/ritik/sih/docs/Thermal-Shelter_SIH2026_PS26051.pptx"
IMG = "/Users/ritik/sih/docs"

ACCENT = RGBColor(0x00, 0x70, 0xC0)   # SIH-style blue
INK    = RGBColor(0x22, 0x2A, 0x33)
GREEN  = RGBColor(0x1D, 0x6B, 0x34)
GREY   = RGBColor(0x66, 0x70, 0x7A)
BODY_F = "Arial"
TITLE_F = "Times New Roman"


# --------------------------------------------------------------------------- #
def style_frame(tf):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE   # "shrink text on overflow"
    tf.margin_left = tf.margin_right = Inches(0.10)
    tf.margin_top = tf.margin_bottom = Inches(0.04)


def fill(tf, rows):
    """rows: list of (text, size, bold, color, level, space_before)."""
    style_frame(tf)
    tf.clear()
    first = tf.paragraphs[0]
    for i, (text, size, bold, color, level, sb) in enumerate(rows):
        p = first if i == 0 else tf.add_paragraph()
        p.level = level
        p.space_before = Pt(sb)
        p.space_after = Pt(2)
        p.line_spacing = 1.02
        r = p.add_run()
        r.text = text
        r.font.name = BODY_F
        r.font.size = Pt(size)
        r.font.bold = bold
        if color is not None:
            r.font.color.rgb = color


def fill_refs(tf, refs):
    """Numbered reference list; each URL becomes a clickable, blue, underlined run.

    refs: list of (description, url_or_None). Renders "1. <desc> — <link>" per row,
    so the references read in sequence and every link is click-through in the PPT/PDF.
    """
    style_frame(tf)
    tf.clear()
    first = tf.paragraphs[0]
    for i, (desc, url) in enumerate(refs):
        p = first if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else 6)
        p.space_after = Pt(2)
        p.line_spacing = 1.02

        num = p.add_run()
        num.text = f"{i + 1}. "
        num.font.name = BODY_F
        num.font.size = Pt(15)
        num.font.bold = True
        num.font.color.rgb = INK

        body = p.add_run()
        body.text = f"{desc} — " if url else desc
        body.font.name = BODY_F
        body.font.size = Pt(15)
        body.font.color.rgb = INK

        if url:
            link = p.add_run()
            link.text = url
            link.font.name = BODY_F
            link.font.size = Pt(15)
            link.hyperlink.address = url          # makes the run click-through
            link.font.color.rgb = ACCENT
            link.font.underline = True


def set_box(shape, left, top, width, height):
    shape.left, shape.top, shape.width, shape.height = (
        Inches(left), Inches(top), Inches(width), Inches(height))


def add_img(slide, name, left, top, width):
    path = os.path.join(IMG, name)
    w, h = Image.open(path).size
    height = width * h / w
    slide.shapes.add_picture(path, Inches(left), Inches(top),
                             Inches(width), Inches(height))
    return top + height


def caption(slide, left, top, width, text):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.5))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.name = BODY_F; r.font.size = Pt(10.5); r.font.italic = True
    r.font.color.rgb = GREY


def find(slide, *name_starts):
    for sh in slide.shapes:
        if sh.name.startswith(name_starts):
            return sh
    return None


# --------------------------------------------------------------------------- #
prs = Presentation(SRC)
S = prs.slides

# ---- Slide 0 : TITLE PAGE ------------------------------------------------- #
title0 = find(S[0], "Title 7")             # template already reads "... HACKATHON 2026"
for para in title0.text_frame.paragraphs:  # defensive: force 2026 even if a template regresses
    for r in para.runs:
        r.text = r.text.replace("2025", "2026")

fill(find(S[0], "TextBox 9").text_frame, [
    ("Problem Statement ID – 26051", 18, True, INK, 0, 0),
    ("Problem Statement Title – Software-Based Model Development for Design of "
     "Area-Specific Shelter for Thermal Comfort Maintenance", 18, True, INK, 0, 8),
    ("Theme – Miscellaneous", 18, True, INK, 0, 8),
    ("PS Category – Software", 18, True, INK, 0, 8),
    ("Team ID – To be allotted after selection", 18, True, INK, 0, 8),
    ("Team Name – HimTaap", 18, True, INK, 0, 8),
])

# ---- Slide 1 : IDEA TITLE + Proposed Solution ----------------------------- #
s1 = S[1]
t1 = find(s1, "Title 1").text_frame
t1.clear()
p = t1.paragraphs[0]
r = p.add_run()
r.text = "Thermal-Shelter — a shelter that stays warm on sunlight, not fuel."
r.font.name = TITLE_F; r.font.size = Pt(25); r.font.bold = True; r.font.color.rgb = ACCENT

body1 = find(s1, "TextBox 8")
set_box(body1, 0.45, 1.45, 7.15, 5.35)
fill(body1.text_frame, [
    ("Proposed solution — what it is", 18, True, ACCENT, 0, 0),
    ("A design tool for cold, high-altitude shelters (Ladakh & border posts) that hold "
     "comfort on sunlight and smart materials — with little or no active heating.", 15, False, INK, 0, 2),
    ("You give it the site climate, the shelter's shape, and its materials; it simulates a "
     "full day–night cycle and outputs exactly what the problem asks:", 15, False, INK, 0, 4),
    ("Indoor temperature over time", 14, False, INK, 1, 0),
    ("Solar thermal energy captured", 14, False, INK, 1, 0),
    ("Heat-flow vs ambient over the period", 14, False, INK, 1, 0),
    ("Then redesign it live — walls, glazing, orientation, thermal mass — and watch the "
     "heating it needs fall toward zero.", 15, False, INK, 0, 4),
    ("What makes it different", 18, True, ACCENT, 0, 8),
    ("Models the night and the cold sky (air −14 °C, sky ≈ −45 °C) — not just the average.", 14, False, INK, 1, 0),
    ("Stores the sun in thermal mass — stone, water walls, PCM.", 14, False, INK, 1, 0),
    ("Area-specific: one engine, any cold site (Leh, Drass, Siachen, Tawang).", 14, False, INK, 1, 0),
    ("Fast — under 1 second per design — so you optimise, not guess. ANSYS validates the winner.", 14, False, INK, 1, 0),
])
bottom = add_img(s1, "demo_curve.png", 7.95, 2.55, 5.05)
caption(s1, 7.95, bottom + 0.05, 5.05,
        "Simulated Leh winter — the passive design (blue) climbs into the comfort "
        "band on sunlight; the baseline hut (red) tracks the freezing air.")

# ---- Slide 2 : TECHNICAL APPROACH ----------------------------------------- #
s2 = S[2]
body2 = find(s2, "TextBox 8")
set_box(body2, 0.45, 1.35, 6.05, 5.4)
fill(body2.text_frame, [
    ("Technologies used", 17, True, ACCENT, 0, 0),
    ("Python 3.12 — NumPy · SciPy · pandas", 14, False, INK, 1, 0),
    ("pvlib — solar position & irradiance: clear-sky or real measured (snow albedo)", 14, False, INK, 1, 0),
    ("PVGIS / EPW — real measured weather (TMY) fetched live for any site", 14, False, INK, 1, 0),
    ("Streamlit + Plotly — the live design dashboard (prototype, right)", 14, False, INK, 1, 0),
    ("matplotlib — figures & reporting", 14, False, INK, 1, 0),
    ("ANSYS — CFD / thermal cross-validation (later phase)", 14, False, INK, 1, 0),
    ("Methodology", 17, True, ACCENT, 0, 8),
    ("Lumped resistance–capacitance (RC) thermal network, stepped through time", 14, False, INK, 1, 0),
    ("Driven by sol-air outdoor temperature + long-wave night-sky cooling", 14, False, INK, 1, 0),
    ("pvlib solar irradiance on every wall, roof and window", 14, False, INK, 1, 0),
    ("Runs on a real measured year (PVGIS TMY / EPW) — rolled up month-by-month", 14, False, INK, 1, 0),
    ("Score comfort (% time in 18–24 °C) and energy (kWh/day to hold 20 °C)", 14, False, INK, 1, 0),
    ("Compare & rank designs by material, size, orientation and glazing", 14, False, INK, 1, 0),
])
add_img(s2, "dashboard.png", 6.8, 1.4, 6.1)
caption(s2, 6.8, 1.4 + 6.1 * 1500 / 2666 + 0.04, 6.1,
        "Working prototype — the live design studio (Streamlit).")
add_img(s2, "deck_flow.png", 6.9, 5.34, 5.8)   # methodology pipeline, right column

# ---- Slide 3 : FEASIBILITY AND VIABILITY ---------------------------------- #
s3 = S[3]
body3 = find(s3, "TextBox 8")
set_box(body3, 0.7, 1.4, 12.0, 5.35)
fill(body3.text_frame, [
    ("Why it's feasible", 17, True, ACCENT, 0, 0),
    ("Built on standard, peer-reviewed physics — ASHRAE sol-air & RC networks, "
     "Berdahl–Martin sky emissivity. Defensible, not a black box.", 14, False, INK, 1, 0),
    ("Uses the mature, widely trusted pvlib solar model.", 14, False, INK, 1, 0),
    ("Pure software — runs on any laptop, with no field hardware to build.", 14, False, INK, 1, 0),
    ("The working prototype already reproduces the passive-design behaviour we expect.", 14, False, INK, 1, 0),
    ("Viability", 17, True, ACCENT, 0, 8),
    ("Cheap, fast and reusable across regions: screen hundreds of designs in software, "
     "then validate only the winner in CFD — where each case would otherwise take hours.", 14, False, INK, 1, 0),
    ("Potential challenges → how we overcome them", 17, True, ACCENT, 0, 8),
    ("Sparse climate data for remote posts → accept user CSV + a typical-day generator "
     "+ built-in region presets.", 14, False, INK, 1, 0),
    ("A lumped RC model simplifies real 3-D heat flow → validate and calibrate the "
     "shortlisted design in ANSYS CFD.", 14, False, INK, 1, 0),
    ("Material properties vary by source → a sourced, editable material-property database.", 14, False, INK, 1, 0),
])

# ---- Slide 4 : IMPACT AND BENEFITS ---------------------------------------- #
s4 = S[4]
body4 = find(s4, "TextBox 8")
set_box(body4, 0.55, 1.4, 6.15, 5.5)
fill(body4.text_frame, [
    ("~93% less heating over a full measured Leh year", 18, True, GREEN, 0, 0),
    ("2,140 vs 31,415 kWh/year to hold 20 °C, from a continuous 8,760-hour "
     "simulation on a PVGIS Typical Meteorological Year (real clouds and irradiance, "
     "not a clear-sky day).", 13, False, INK, 0, 3),
    ("≈ 3,550 L kerosene · ₹3.2 lakh · ~8.9 t CO₂ — saved per shelter, per year.", 15, True, ACCENT, 0, 5),
    ("(Clear design-day: 82–86% less across Leh, Drass, Siachen, Tawang — Leh 84%.)", 12, False, GREY, 0, 4),
    ("Who it helps", 16, True, ACCENT, 0, 8),
    ("Armed forces & border posts · high-altitude communities · disaster relief.", 13, False, INK, 1, 0),
    ("Benefits", 16, True, ACCENT, 0, 8),
    ("Economic — ~3,550 fewer litres of kerosene and ₹3.2 lakh saved per shelter each "
     "year; designs done in minutes, not months.", 13, False, INK, 1, 0),
    ("Environmental — ~8.9 t less CO₂ per shelter per year in a fragile Himalayan ecosystem.", 13, False, INK, 1, 0),
    ("Operational & social — energy resilience where supply lines are thin; warmer, safer "
     "shelters; fewer cold-related injuries.", 13, False, INK, 1, 0),
])
bottom = add_img(s4, "seasonal.png", 6.9, 1.55, 6.05)
caption(s4, 6.9, bottom + 0.05, 6.05,
        "All year on real measured weather — monthly heating (design vs baseline hut), "
        "% time comfortable, and the fuel / ₹ / CO₂ saved per shelter per year.")

# ---- Slide 5 : RESEARCH AND REFERENCES ------------------------------------ #
s5 = S[5]
body5 = find(s5, "TextBox 8")
set_box(body5, 0.7, 1.5, 12.0, 5.2)
fill_refs(body5.text_frame, [
    ("pvlib-python (Sandia National Laboratories) — solar position & clear-sky "
     "irradiance", "https://pvlib-python.readthedocs.io"),
    ("PVGIS (EU Joint Research Centre) — Typical Meteorological Year, free, no API key",
     "https://re.jrc.ec.europa.eu/pvg_tools/en/"),
    ("Berdahl, P. & Martin, M. (1984), “Emissivity of clear skies” — long-wave "
     "radiative night-sky cooling", None),
    ("ASHRAE Handbook of Fundamentals — sol-air temperature, RC thermal networks, "
     "U-values & SHGC", None),
    ("Balcomb, J. D. / U.S. Department of Energy — passive-solar design: Trombe & "
     "water walls, thermal mass", None),
    ("DRDO–DIHAR (Defence Institute of High-Altitude Research), Leh — high-altitude "
     "habitat & thermal-comfort context", None),
    ("Project repository", "https://github.com/ritikx-36/sih-2026"),
])

# ---- team-name ovals on every content slide ------------------------------- #
for s in (S[1], S[2], S[3], S[4], S[5]):
    ov = find(s, "Oval")
    if ov is not None and ov.has_text_frame:
        ov.text_frame.paragraphs[0].runs and setattr(
            ov.text_frame.paragraphs[0].runs[0], "text", "HimTaap")

# ---- delete the "Important Instructions" slide (index 6) ------------------ #
sldIdLst = prs.slides._sldIdLst
sldIdLst.remove(list(sldIdLst)[6])

prs.save(OUT)
print(f"saved -> {OUT}  ({len(prs.slides)} slides)")
